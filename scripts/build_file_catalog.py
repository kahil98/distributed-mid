from collections import defaultdict
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Distributed_Storage_Project_File_Catalog.docx"
SKIP_DIRS = {".git", ".venv", "node_modules", "__pycache__", ".pytest_cache", "build-wheels", "dist"}

def include(path):
    rel = path.relative_to(ROOT)
    if any(part in SKIP_DIRS for part in rel.parts): return False
    if path.name == ".env" or (path.name.startswith(".env.") and path.name != ".env.example"): return False
    if path.suffix in {".pyc", ".pyo"}: return False
    if path.name.endswith("-submission.zip"): return False
    return True

def describe(path):
    name, parent, suffix = path.name, path.parent.name, path.suffix.lower()
    known = {
        "app.py": "FastAPI application entry point and service routes.",
        "chunking.py": "Fixed-size and content-defined chunking algorithms.",
        "ring.py": "Consistent-hash ring and node placement helpers.",
        "reliability.py": "Replication and erasure-coding policies and operations.",
        "job_worker.py": "RabbitMQ maintenance-job consumer and dispatcher.",
        "verify_containers.py": "Live verification of storage-node containers and chunk operations.",
        "verify_metadata.py": "Live verification of metadata, chunking, deduplication, and placement.",
        "verify_reliability.py": "Live verification of durability, repair, and scrub behavior.",
        "verify_gateway.py": "Live verification of the authenticated gateway API.",
        "benchmark_storage.py": "Generates and verifies synthetic small-file and large-file workloads.",
        "build_project_report.py": "Source generator for the project implementation report.",
        "build_file_catalog.py": "Source generator for this project file catalog.",
        "package_submission.py": "Packages the project source and evidence while omitting secrets and caches.",
        "test_storage_node.py": "Automated tests for the storage-node HTTP API and disk integrity.",
        "test_metadata_algorithms.py": "Automated tests for chunking and consistent-hash placement.",
        "test_reliability.py": "Automated tests for durability policies and erasure coding.",
        "test_gateway.py": "Automated tests for gateway authentication and rate limiting.",
        "test_task6_jobs.py": "Automated tests for publishing and dispatching maintenance jobs.",
        "nginx.conf": "Nginx reverse-proxy configuration for gateway traffic.",
        "prometheus.yml": "Prometheus scrape targets and intervals.",
        "distributed-storage.json": "Provisioned Grafana dashboard panels and Prometheus queries.",
        "task8-cloud.yaml": "Provider-neutral Kubernetes gateway, autoscaling, load balancer, and storage template.",
        "storage-nodes.yaml": "Kubernetes storage-node StatefulSet, headless service, PVCs, and probes.",
        "compose.yaml": "Local Prometheus, Grafana, Kafka, and RabbitMQ services.",
        "compose.storage.yaml": "Local metadata, PostgreSQL, Elasticsearch, gateway, Nginx, workers, nodes, and UI stack.",
        "Dockerfile": "Container build recipe for the associated service or root application image.",
        "Dockerfile.offline": "Offline/fallback application image build recipe.",
        "package.json": "Frontend package metadata and npm scripts/dependencies.",
        "package-lock.json": "Pinned frontend npm dependency graph.",
        "README.md": "Project or task setup, implementation, testing, and scope documentation.",
        "SUBMISSION.md": "Contents and quick verification instructions for the submitted source package.",
    }
    if name in known: return known[name]
    if parent == "reports": return "Saved test, runtime verification, or benchmark evidence."
    if parent == "tests": return "Automated pytest coverage for the project."
    if parent == "docs": return "Project task, architecture, or setup documentation."
    if parent == "infra": return "Infrastructure or observability configuration."
    if parent == "postman": return "Postman collection for API demonstration and testing."
    if parent == "src": return "React frontend source code or styling."
    if parent == "public": return "Static frontend asset served by Vite."
    if suffix == ".py": return "Python implementation or utility script."
    if suffix in {".yaml", ".yml"}: return "YAML service, deployment, or tool configuration."
    if suffix == ".json": return "Structured project, test, or tool data."
    if suffix in {".txt", ".lock"}: return "Dependency, lock, or runtime configuration file."
    if suffix in {".css", ".jsx", ".js"}: return "Frontend application source or build configuration."
    if suffix == ".svg": return "Scalable vector frontend asset."
    return "Project source, documentation, or configuration asset."

files = sorted((p for p in ROOT.rglob("*") if p.is_file() and include(p)), key=lambda p: p.relative_to(ROOT).as_posix().lower())
groups = defaultdict(list)
for p in files:
    rel = p.relative_to(ROOT)
    group = rel.parts[0] if len(rel.parts) > 1 else "Root files"
    groups[group].append(p)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(9)
doc.styles["Title"].font.name = "Aptos Display"; doc.styles["Title"].font.size = Pt(25); doc.styles["Title"].font.color.rgb = RGBColor(31,78,121)
for s in ("Heading 1", "Heading 2"):
    doc.styles[s].font.name = "Aptos Display"; doc.styles[s].font.color.rgb = RGBColor(31,78,121)
doc.styles["Heading 1"].font.size = Pt(15); doc.styles["Heading 2"].font.size = Pt(11)
doc.add_paragraph("Project File Catalog", style="Title")
doc.add_paragraph("Distributed object storage framework | Complete source and evidence inventory")
doc.add_paragraph(f"Cataloged {len(files)} files from the project workspace. Local secrets and generated dependency/cache folders are excluded.")
doc.add_heading("How to read this catalog", level=1)
doc.add_paragraph("Entries are grouped by their top-level directory. Each entry gives the repository-relative path and its role. This catalog covers the portable project deliverable; local .env secrets, .venv, node_modules, caches, and generated build output are excluded intentionally.")
doc.add_heading("File inventory", level=1)
for group in sorted(groups, key=lambda s: (s != "Root files", s.lower())):
    doc.add_heading(group + ("/" if group != "Root files" else ""), level=2)
    for path in groups[group]:
        rel = path.relative_to(ROOT).as_posix()
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(rel + " — "); run.bold = True; run.font.name = "Consolas"; run.font.size = Pt(8)
        p.add_run(describe(path))
doc.add_heading("Related documents", level=1)
doc.add_paragraph("The implementation narrative is in Distributed_Storage_Project_Full_Report.docx. The technical architecture overview is in docs/system-architecture.md. Task-specific startup and test procedures are linked from README.md.")
doc.core_properties.title = "Distributed Storage Project File Catalog"
doc.core_properties.subject = "Inventory and purpose of project deliverable files"
doc.core_properties.author = "Course Project Team"
doc.save(OUT)
print(f"{OUT}\nfiles={len(files)} groups={len(groups)}")
