from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"Architecture_Test_and_Benchmark_Report.docx"
d=Document(); sec=d.sections[0]
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
d.styles["Normal"].font.name="Aptos"; d.styles["Normal"].font.size=Pt(9.5)
d.styles["Title"].font.name="Aptos Display"; d.styles["Title"].font.size=Pt(25); d.styles["Title"].font.color.rgb=RGBColor(31,78,121)
for style in ("Heading 1","Heading 2"):
    d.styles[style].font.name="Aptos Display"; d.styles[style].font.color.rgb=RGBColor(31,78,121)
d.styles["Heading 1"].font.size=Pt(16); d.styles["Heading 2"].font.size=Pt(11)

def bullet(label, text):
    p=d.add_paragraph(style="List Bullet"); p.add_run(label+" ").bold=True; p.add_run(text)
def command(text):
    p=d.add_paragraph(); p.paragraph_format.left_indent=Inches(.22); r=p.add_run(text); r.font.name="Consolas"; r.font.size=Pt(8)

d.add_paragraph("System Architecture, Test Scripts, and Benchmark Results",style="Title")
d.add_paragraph("Distributed Object Storage Framework | Course Project Report")
d.add_paragraph("Prepared 29 October 2026")
d.add_heading("Purpose and scope",level=1)
d.add_paragraph("This report documents how the project is organized, how requests and data move through the system, which verification tools are supplied, and what the saved workload and reliability results show. Benchmark numbers are quoted from the repository reports and include their workload and measurement limitations. This report describes the implementation artifacts; it does not claim that the most recent Docker or cloud deployment completed successfully.")

d.add_heading("1. System architecture",level=1)
d.add_paragraph("The system separates the public API, metadata/control plane, and chunk data plane. A browser uses the React/Vite console, while scripts and other clients can use the HTTP API directly. Nginx is intended to receive gateway traffic and proxy it to multiple FastAPI gateway replicas. The gateway handles login, JWT authorization, rate limiting, object and bucket endpoints, search, and cluster summaries. It delegates storage operations to the metadata service.")
d.add_paragraph("The metadata service selects fixed-size or content-defined chunks, computes SHA-256 addresses, maps chunks to storage nodes with a virtual-node consistent-hash ring, and records object manifests and placement state in PostgreSQL. Depending on bucket policy, data is stored as replicas or Reed-Solomon fragments. Storage nodes persist chunk files on node-local disks or Kubernetes persistent volumes and periodically send capacity and chunk-count heartbeats. Metadata schedules repair, rebalance, and garbage-collection work through RabbitMQ; workers invoke maintenance handlers. Elasticsearch backs object metadata search. Prometheus scrapes metrics and Grafana displays them.")
d.add_heading("Request flow",level=2)
for x in [
    "Upload: client -> Nginx -> gateway (JWT and request checks) -> metadata service (chunking and policy) -> storage nodes (chunk or fragment PUT). PostgreSQL records the manifest and references; Elasticsearch indexes metadata.",
    "Download: client -> gateway -> metadata service -> enough healthy storage fragments. The metadata service reconstructs erasure-coded data when needed and verifies the logical digest before returning it; the gateway can serve a byte range.",
    "Failure and maintenance: heartbeat timeout or scrub corruption updates placement state and creates repair work. RabbitMQ distributes jobs to worker processes, which call metadata repair/rebalance/GC handlers.",
    "Monitoring: Prometheus scrapes gateway request/latency metrics, storage-node free bytes/chunk count, and metadata degraded-object/repair counters; Grafana queries Prometheus."]:
    d.add_paragraph(x,style="List Bullet")
d.add_paragraph("Local endpoints used by the project include Nginx/gateway at 127.0.0.1:8080, metadata at 127.0.0.1:8100, PostgreSQL at 127.0.0.1:5433, Elasticsearch at 127.0.0.1:9200, storage nodes at ports 8201 onward, Prometheus at 127.0.0.1:9090, Grafana at 127.0.0.1:3000, and RabbitMQ management at 127.0.0.1:15672.")

d.add_heading("2. Project structure",level=1)
for label,text in [
    ("gateway_service/","Authenticated user API, JWT, request limiting, streaming, range downloads, object search and cluster view."),
    ("metadata_service/","Object manifests, PostgreSQL schema, fixed/CDC chunking, consistent-hash ring, replication/erasure coding, repair and scrubbing."),
    ("storage_node/","Chunk HTTP operations, disk persistence, SHA-256 validation, health and heartbeat metrics."),
    ("frontend/","React/Vite interface, npm dependency lock, styles, and UI container recipe."),
    ("scripts/","Live verification programs, synthetic workload benchmark, RabbitMQ worker, and report/package helpers."),
    ("tests/","Pytest checks for storage-node behavior, chunking/placement algorithms, reliability, gateway auth, and job routing."),
    ("compose.yaml and compose.storage.yaml","Local infrastructure and application service definitions."),
    ("deploy/kubernetes/","Storage StatefulSet/PVC/probe definitions and provider-neutral gateway HPA/LoadBalancer template."),
    ("infra/","Nginx reverse proxy, Prometheus scrape configuration, and Grafana datasource/dashboard provisioning."),
    ("reports/","Saved JSON/XML unit, integration, reliability, gateway, and benchmark outcomes."),
    ("docs/ and postman/","Task setup/verification guides, architecture report, and a Postman gateway collection.")]: bullet(label,text)

d.add_heading("3. Verification scripts and automated tests",level=1)
d.add_paragraph("The pytest suite is the fast regression layer. The scripts below perform live checks against services, so Compose containers must be running and healthy first.")
for label,text in [
    ("tests/test_storage_node.py","Exercises PUT/GET/HEAD/DELETE, invalid digests, mismatched content hashes, corruption detection, and disk behavior in isolated tests."),
    ("tests/test_metadata_algorithms.py","Checks deterministic ring placement, movement expectations, fixed/CDC chunk boundaries, and reconstruction properties."),
    ("tests/test_reliability.py","Checks durability policy validation and Reed-Solomon encode/decode behavior."),
    ("tests/test_gateway.py","Checks JWT login/authorization and rate-limit responses."),
    ("tests/test_task6_jobs.py","Checks RabbitMQ job publisher routing and worker handler selection."),
    ("scripts/verify_containers.py","Checks four live storage containers, heartbeats, operations, persistence across restart, corruption rejection, re-upload recovery, and independent node disks. Saved output: reports/container-verification.json."),
    ("scripts/verify_metadata.py","Checks metadata/PostgreSQL health, manifests, deduplication/refcounts, object versions, fixed and CDC transfer, and ring movement. Saved output: reports/task-2-verification.json."),
    ("scripts/verify_reliability.py","Checks quorum policy, replication and 4+2 recoverability, node loss, corruption/scrub repair, and bounded storage/repair traffic comparisons. Saved output: reports/task-3-verification.json."),
    ("scripts/verify_gateway.py","Checks JWT, buckets, uploads, tags/search, range/full download, cluster status, deletion, and Elasticsearch removal. Saved output: reports/task-4-verification.json."),
    ("scripts/benchmark_storage.py","Generates deterministic synthetic data, sends chunks to storage-node endpoints, verifies each chunk and reassembled file hash, and records latency and elapsed throughput. Saved output: reports/workloads.json.")]: bullet(label,text)
d.add_paragraph("Most recent automated run recorded during report preparation: 54 tests passed. Compose configuration validation passed. Live service scripts should be run after a successful image build and healthy Compose startup.")
d.add_heading("Re-run commands",level=2)
for c in [r".\.venv\Scripts\python.exe -m pytest -q",r"docker compose -f compose.storage.yaml config --quiet",r".\.venv\Scripts\python.exe scripts/verify_containers.py",r".\.venv\Scripts\python.exe scripts/verify_metadata.py",r".\.venv\Scripts\python.exe scripts/verify_reliability.py",r".\.venv\Scripts\python.exe scripts/verify_gateway.py",r".\.venv\Scripts\python.exe scripts/benchmark_storage.py"]: command(c)

d.add_heading("4. Benchmark results",level=1)
d.add_paragraph("The Task 1 workload benchmark targets the storage-node chunk API. It performs fixed-size 4 MiB splitting and client-side reassembly, not the later metadata coordinator, replication, erasure coding, or S3 interface. Throughput includes data generation and cleanup. Every recorded file and chunk hash passed verification.")
d.add_heading("Many small files",level=2)
d.add_paragraph("The saved run processed 5,000 files ranging from 1–50 KB. Total input was 131,867,328 bytes across 5,000 chunks, using eight workers. Elapsed time was 34.199 seconds and end-to-end throughput including generation and cleanup was 7.355 MiB/s. PUT latency was 27.688 ms median and 54.125 ms p95. GET latency was 9.948 ms median and 19.660 ms p95. All file and chunk hashes matched.")
d.add_heading("Few large files",level=2)
d.add_paragraph("The saved run processed 10 files of 100 MiB each, totaling 1,048,576,000 bytes and 250 chunks, using four workers. Elapsed time was 36.715 seconds and end-to-end throughput including generation and cleanup was 54.473 MiB/s. PUT latency was 318.571 ms median and 497.421 ms p95. GET latency was 150.082 ms median and 228.639 ms p95. All file and chunk hashes matched.")
d.add_heading("Placement movement and deduplication",level=2)
d.add_paragraph("Task 2 measured 50,000 generated hashes on a ring with 128 virtual nodes. Adding a fifth node moved 10,801 hashes (21.602%), compared with an approximate 20% expectation. Removing one of four nodes moved 12,955 hashes (25.91%), compared with an approximate 25% expectation. The metadata verification also recorded successful deduplication: two objects sharing a 4 MiB chunk produced a reference count of two; deleting one object kept the chunk, and deleting the final reference removed it.")
d.add_heading("Replication versus erasure coding samples",level=2)
d.add_paragraph("The Task 3 report compares three-copy replication with 4+2 Reed-Solomon coding. In a 12-file, 49,900-byte sample, replication stored 149,700 bytes (3.00x) and estimated one-node repair traffic at 33,266 bytes; erasure coding stored 74,874 bytes (1.5005x) and estimated repair traffic at 62,375 bytes. In a two-file, 163,840-byte sample, replication stored 491,520 bytes (3.00x) with estimated repair traffic of 109,226 bytes; erasure coding stored 245,760 bytes (1.50x) with estimated repair traffic of 204,800 bytes.")
d.add_paragraph("These are bounded samples, not the full Task 1 workload sizes. Erasure coding used about half the physical storage in these samples, while one-fragment repair required reading multiple shares and therefore had higher estimated traffic. The same saved reliability run reports 11,092,880 measured repair bytes, 37,386,252 scrub-read bytes, and 13 completed repair jobs across its failure/scrub checks. Estimated per-sample traffic and measured service counters are different measurements and should not be combined.")

d.add_heading("5. Evidence files and interpretation",level=1)
for label,text in [
    ("reports/workloads.json","Raw Task 1 workload timing, size, chunk count, throughput, latency, and hash-verification results."),
    ("reports/container-verification.json","35 recorded live storage-node checks."),
    ("reports/task-2-verification.json","17 recorded metadata and placement checks, including the 50,000-hash add/remove simulation."),
    ("reports/task-3-verification.json","Recorded reliability checks, measured service repair/scrub counters, and bounded policy comparison samples."),
    ("reports/task-4-verification.json","Gateway workflow and search verification results."),
    ("reports/unit-tests.xml and reports/task-2-unit-tests.xml","JUnit-style automated test result records from prior runs.")]: bullet(label,text)
d.add_paragraph("Results are snapshots from saved runs, not a guarantee of current live service health. The latest Docker build attempt reported an `exec format error`, and the cloud manifest has not been deployed to a cloud provider. Local PostgreSQL is single-instance. Task 7 requirements were not present in the supplied project materials.")
d.add_heading("Conclusion",level=1)
d.add_paragraph("The repository documents a full development path from verified chunk storage through metadata placement, durability, user APIs, a browser console, background operations, and deployment templates. The strongest quantitative evidence is the verified Task 1 workload run and the saved Task 2/3 placement and durability reports. For an exam demonstration, rerun the live scripts after Docker builds successfully, capture current Prometheus/Grafana and UI evidence, and describe the cloud Kubernetes files as templates unless an actual cloud cluster is deployed.")
d.core_properties.title="System Architecture Test and Benchmark Report"
d.core_properties.subject="Architecture, repository structure, verification scripts, and benchmark results"
d.core_properties.author="Course Project Team"
d.save(OUT)
print(OUT)
